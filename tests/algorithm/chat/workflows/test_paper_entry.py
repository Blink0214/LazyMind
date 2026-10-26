"""Direct manuscript entry: real files and readers, without a model replacement."""
import json

import pytest

from lazymind.document_tools.paper import workflow
from lazymind.document_tools.paper.entry import read_manuscript


@pytest.mark.parametrize('command', [
    '重试', '重试！', 'retry', '请重新执行步骤 peer_review', 'Please re-run step peer_review',
])
def test_retry_control_preserves_review_requirements(command):
    params = {'current_user_input': command, 'user_input': '重点检查数字与实验设计'}
    request, _ = workflow._runtime_request(params)
    pool = workflow.initialize_pool({}, request)
    assert workflow._step_request('peer_review', pool, params) == '重点检查数字与实验设计'


def test_retry_uses_saved_module_requirements_not_initial_launch():
    pool = workflow.pool_from_business({'topic': 'T', 'request': '检查论文', 'manuscript': '正文'})
    pool['history'] = [{'node': 'F6', 'request': '只检查引用'}, {'node': 'F5', 'request': '修改措辞'}]
    params = {'retry_hint': 'Recovery request for this rerun only: 重试', 'user_input': '检查论文'}
    assert workflow._step_request('peer_review', pool, params) == '只检查引用'
    assert pool['request'] == '检查论文'


@pytest.mark.parametrize('params', [
    {'current_user_input': '重试，只检查引用'},
    {'retry_hint': 'Recovery request for this rerun only: 只检查引用'},
])
def test_explicit_new_requirements_override_saved_request(params):
    pool = workflow.pool_from_business({'topic': 'T', 'request': '检查数字', 'manuscript': '正文'})
    assert '只检查引用' in workflow._step_request('peer_review', pool, params)


def test_retry_does_not_replace_bound_review_request():
    pool = workflow.pool_from_business({'topic': 'T', 'request': '检查可复现性', 'manuscript': '正文'})
    params = {'current_user_input': '重试', 'launch_user_input': '审稿'}
    assert workflow._step_request('peer_review', pool, params) == '检查可复现性'


@pytest.mark.parametrize('step', ['refine_topic', 'experiment', 'peer_review'])
def test_forward_step_uses_current_user_input_from_runtime(step):
    pool = workflow.pool_from_business({'topic': 'T', 'request': '最初只做选题'})
    pool['history'] = [{'node': 'F2', 'request': '上轮综述要求'}]
    params = {'step_id': step, 'user_input': '最初只做选题',
              'current_user_input': '综述之后调整选题，只研究失败恢复，不做实验',
              'remote_inputs': {'literature_review_pool': '/saved/pool.json'}}
    assert workflow._step_request(step, pool, params) == params['current_user_input']
    assert pool['request'] == '最初只做选题'


def test_retry_uses_saved_refinement_requirements():
    pool = workflow.pool_from_business({'topic': 'T', 'request': '最初选题'})
    pool['history'] = [{'node': 'F1', 'request': '聚焦失败恢复'}]
    assert workflow._step_request('refine_topic', pool, {'current_user_input': '重试'}) == '聚焦失败恢复'


def test_launch_intent_does_not_override_explicit_bound_requirements(tmp_path):
    pool = workflow.pool_from_business({'topic': 'T', 'request': '只检查引用'})
    request_file = tmp_path / 'review-request.txt'
    request_file.write_text('只检查引用')
    params = {'step_id': 'peer_review', 'user_input': '审稿', 'launch_user_input': '审稿',
              'current_user_input': ''}
    assert workflow._step_request('peer_review', pool, params,
                                  {'review_request': str(request_file)}) == '只检查引用'


def test_recovery_request_wins_over_restored_launch_input():
    pool = workflow.pool_from_business({'topic': 'T', 'request': '最初选题'})
    pool['history'] = [{'node': 'F2', 'request': '旧综述'}]
    params = {'user_input': '最初选题', 'retry_hint': 'Recovery request for this rerun only: 只调整恢复机制',
              'remote_inputs': {'literature_review_pool': '/saved/pool.json'}}
    assert workflow._step_request('refine_topic', pool, params) == '只调整恢复机制'


def test_text_upload_initializes_without_pool_or_json(tmp_path):
    source = tmp_path / 'paper.md'
    source.write_text('# My paper\n## Results\n12 of 20 questions were correct.')
    pool = workflow.initialize_pool({'manuscript_file': str(source)}, 'Review this paper')
    assert pool['topic'] == 'My paper'
    assert workflow.material(pool, 'manuscript')['content'] == source.read_text()
    assert pool['history'] == []
    assert not workflow.material(pool, 'manuscript')['mock']
    workflow.contracts.validate_input('peer_review', workflow.business_input('peer_review', pool))


def test_plaintext_body_and_requirements_without_business_json(tmp_path):
    body = tmp_path / 'body'
    body.write_text('# Title\nFull manuscript')
    request = tmp_path / 'request'
    request.write_text('Only check evidence')
    pool = workflow.initialize_pool({'manuscript_text': str(body), 'review_request': str(request)}, 'launch command')
    # Review requirements are resolved by the executing step, not mixed into materials.
    assert pool['request'] == 'launch command'
    assert workflow._step_request('peer_review', pool, {}, {'review_request': str(request)}) == 'Only check evidence'
    assert pool['topic'] == 'Title'


def test_latest_pool_wins_over_original_upload(tmp_path):
    pool = workflow.pool_from_business({'topic': 'T', 'manuscript': 'Latest approved manuscript'})
    saved = tmp_path / 'saved.json'
    saved.write_text(json.dumps(pool))
    result = workflow.initialize_pool({'write_document_pool': str(saved), 'manuscript_file': '/missing/original.pdf'})
    assert workflow.material(result, 'manuscript')['content'] == 'Latest approved manuscript'


def test_conflicting_manuscript_sources_rejected(tmp_path):
    source = tmp_path / 'paper.md'
    source.write_text('# Manuscript')
    with pytest.raises(ValueError, match='同时提供'):
        workflow.initialize_pool({'manuscript_file': str(source), 'manuscript_text': str(source)})


def test_missing_manuscript_cannot_enter_review():
    pool = workflow.initialize_pool({}, 'Review this paper')
    with pytest.raises(workflow.MissingMaterials, match='manuscript'):
        workflow.contracts.validate_input('peer_review', workflow.business_input('peer_review', pool))


@pytest.mark.parametrize('name,content', [('empty.txt', ''), ('paper.exe', 'binary')])
def test_unusable_upload_is_rejected(tmp_path, name, content):
    source = tmp_path / name
    source.write_text(content)
    with pytest.raises(ValueError):
        read_manuscript(str(source))


def test_docx_reader_preserves_body(tmp_path):
    from docx import Document
    source = tmp_path / 'paper.docx'
    document = Document()
    document.add_heading('My manuscript', 0)
    document.add_paragraph('The method answered 15 of 20 questions correctly.')
    document.save(source)
    content = read_manuscript(str(source))
    assert 'My manuscript' in content
    assert '15 of 20' in content


def test_pdf_reader_preserves_body(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    source = tmp_path / 'paper.pdf'
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({
        NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (My paper: 15 of 20 questions correct.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    with source.open('wb') as output:
        writer.write(output)
    assert '15 of 20' in read_manuscript(str(source))


def test_image_only_pdf_does_not_silently_review_empty_text(tmp_path):
    from pypdf import PdfWriter
    source = tmp_path / 'scan.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with source.open('wb') as output:
        writer.write(output)
    with pytest.raises(ValueError, match='提取文字'):
        read_manuscript(str(source))
